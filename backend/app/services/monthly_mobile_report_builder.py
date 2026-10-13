import calendar
import io
import re
from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.daily_hour_processor import HOURS
from app.services.report_session_store import report_session_store


FONT_NAME = "Calibri"

THIN_SIDE = Side(style="thin", color="000000")
DOUBLE_SIDE = Side(style="double", color="000000")
MEDIUM_SIDE = Side(style="medium", color="000000")

THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
HEADER_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
TOTAL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=DOUBLE_SIDE)

SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")

HOURLY_DISPLAY_HEADERS = [
    "00:00 - 01:00",
    "01:00 - 02:00",
    "02:00 - 03:00",
    "03:00 - 04:00",
    "04:00 - 05:00",
    "05:00 - 06:00",
    "06:00 - 07:00",
    "07:00 - 08:00",
    "08:00 - 09:00",
    "09:00 - 10:00",
    "10:00 - 11:00",
    "11:00 - 12:00",
    "12:00 - 13:00",
    "13:00 - 14:00",
    "14:00 - 15:00",
    "15:00 - 16:00",
    "16:00 - 17:00",
    "17:00 - 18:00",
    "18:00 - 19:00",
    "19:00 - 20:00",
    "20:00 - 21:00",
    "21:00 - 22:00",
    "22:00 - 23:00",
    "23:00 - 00:00",
]

COMPLIANCE_HEADERS = [
    "Date",
    "Total Weighed (X)",
    "Warned & Charged (Y)",
    "Charged (Z)",
    "Warned (A)",
    "Cases Cleared in Court (B)",
    "  Transgressions (L)",
    "Exempted permit (E)",
    "Manually Weighed (M)",
]


def _safe_int(val: Any) -> int:
    try:
        return int(float(str(val).replace(",", "").strip()))
    except (ValueError, TypeError, AttributeError):
        return 0


def _mobile_report_slot(bound_name: str | None) -> str:
    bound = (bound_name or "").strip().lower()
    if "2" in bound or "two" in bound:
        return "mobile_2"
    return "mobile_1"


def _extract_session_stats(session) -> dict[str, Any]:
    summary = session.sections.get("mobile_report", {}).get("summary", {})
    if not summary and hasattr(session, "dataframes") and "mobile_report" in session.dataframes:
        from app.services.mobile_report_processor import summarize_mobile_report
        df = session.dataframes["mobile_report"]
        if isinstance(df, pd.DataFrame) and not df.empty:
            summary = summarize_mobile_report(df)

    weighed = _safe_int(summary.get("total_trucks_weighed", 0))
    warned = _safe_int(summary.get("warned_trucks", 0))
    charged = _safe_int(summary.get("charged_trucks", 0))

    manual = session.manual_inputs.get("mobile_report") or session.manual_inputs.get("extra", {}).get("mobile_report") or {}
    cases_cleared = _safe_int(manual.get("cases_cleared_in_court", 0))
    transgressions = _safe_int(manual.get("transgressions_count", 0))
    exempted = _safe_int(manual.get("exempted_permit", 0))
    manually_weighed = _safe_int(manual.get("manually_weighed", 0))

    raw_hourly = summary.get("hourly_counts", {})
    hourly = {h: _safe_int(raw_hourly.get(h, 0)) for h in HOURS}

    return {
        "weighed": weighed,
        "warned_charged": warned + charged,
        "charged": charged,
        "warned": warned,
        "cases_cleared": cases_cleared,
        "transgressions": transgressions,
        "exempted": exempted,
        "manually_weighed": manually_weighed,
        "hourly": hourly,
    }


def collect_monthly_mobile_data(station: str, year: int, month: int) -> dict[str, Any]:
    _, num_days = calendar.monthrange(year, month)
    dates = [f"{year:04d}-{month:02d}-{day:02d}" for day in range(1, num_days + 1)]
    dates_set = set(dates)

    norm_station = re.sub(r"[^a-z0-9]", "", station.lower()).replace("mobile", "")

    # Query all valid mobile sessions in the store
    matching_sessions: list[tuple[Any, Any]] = []
    for s, modified_at in report_session_store.list_all_sessions():
        if not s or not s.report_date or s.report_date not in dates_set:
            continue

        sess_station = (s.station or s.weighbridge_name or "").lower()
        sess_bound = (s.bound or "").lower()
        has_mobile_ready = s.sections.get("mobile_report", {}).get("status") == "ready"
        is_mobile_bound = "mobile" in sess_bound

        if not (has_mobile_ready or is_mobile_bound):
            continue

        sess_norm = re.sub(r"[^a-z0-9]", "", sess_station).replace("mobile", "")
        if norm_station in sess_norm or sess_norm in norm_station:
            matching_sessions.append((s, modified_at or 0))

    shift_a_by_date: dict[str, dict[str, Any]] = {}
    shift_b_by_date: dict[str, dict[str, Any]] = {}

    for d in dates:
        cand_a = [
            (s, mod) for s, mod in matching_sessions
            if s.report_date == d and _mobile_report_slot(s.bound) == "mobile_1"
        ]
        cand_a.sort(
            key=lambda item: (item[0].sections.get("mobile_report", {}).get("status") == "ready", item[1]),
            reverse=True,
        )
        s_a = cand_a[0][0] if cand_a else None

        cand_b = [
            (s, mod) for s, mod in matching_sessions
            if s.report_date == d and _mobile_report_slot(s.bound) == "mobile_2"
        ]
        cand_b.sort(
            key=lambda item: (item[0].sections.get("mobile_report", {}).get("status") == "ready", item[1]),
            reverse=True,
        )
        s_b = cand_b[0][0] if cand_b else None

        shift_a_by_date[d] = _extract_session_stats(s_a) if s_a else {
            "weighed": 0,
            "warned_charged": 0,
            "charged": 0,
            "warned": 0,
            "cases_cleared": 0,
            "transgressions": 0,
            "exempted": 0,
            "manually_weighed": 0,
            "hourly": {h: 0 for h in HOURS},
        }

        shift_b_by_date[d] = _extract_session_stats(s_b) if s_b else {
            "weighed": 0,
            "warned_charged": 0,
            "charged": 0,
            "warned": 0,
            "cases_cleared": 0,
            "transgressions": 0,
            "exempted": 0,
            "manually_weighed": 0,
            "hourly": {h: 0 for h in HOURS},
        }

    return {
        "dates": dates,
        "num_days": num_days,
        "shift_a": shift_a_by_date,
        "shift_b": shift_b_by_date,
    }


def apply_outer_border(
    ws,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    outer_side=MEDIUM_SIDE,
    inner_side=THIN_SIDE,
):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            top = outer_side if r == min_row else inner_side
            bottom = outer_side if r == max_row else inner_side
            left = outer_side if c == min_col else inner_side
            right = outer_side if c == max_col else inner_side
            ws.cell(row=r, column=c).border = Border(
                left=left, right=right, top=top, bottom=bottom
            )


def build_monthly_mobile_excel_report(station: str, year: int, month: int) -> io.BytesIO:
    data = collect_monthly_mobile_data(station, year, month)
    dates = data["dates"]
    num_days = data["num_days"]
    shift_a = data["shift_a"]
    shift_b = data["shift_b"]

    wb = Workbook()

    # Fonts
    font_title = Font(name=FONT_NAME, size=11, bold=True)
    font_bold = Font(name=FONT_NAME, size=10, bold=True)
    font_regular = Font(name=FONT_NAME, size=10)
    font_header_8 = Font(name=FONT_NAME, size=8, bold=True)
    font_date_11 = Font(name=FONT_NAME, size=11, bold=True)
    font_data_11 = Font(name=FONT_NAME, size=11, bold=False)
    font_summary_11 = Font(name=FONT_NAME, size=11, bold=True)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_vertical_header = Alignment(horizontal="center", vertical="center", textRotation=90)
    align_wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_middle_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # =========================================================================
    # SHEET 1: MOBILE SUMMARY {YEAR}
    # =========================================================================
    ws1 = wb.active
    ws1.title = f"MOBILE SUMMARY  {year}"
    ws1.views.sheetView[0].showGridLines = True

    current_row = 1

    LIGHT_GRAY_5_FILL = PatternFill("solid", fgColor="F2F2F2")

    def write_compliance_block(title: str, dataset_type: str):
        nonlocal current_row

        # Title Row: merged B to J
        ws1.row_dimensions[current_row].height = 20
        ws1.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=10)
        t_cell = ws1.cell(row=current_row, column=2, value=title)
        t_cell.font = font_title
        t_cell.alignment = align_center

        for col in range(2, 11):
            ws1.cell(row=current_row, column=col).border = THIN_BORDER
        current_row += 1

        # Header Row (Row 2, B2:J2) - text rotated bottom to top (90 deg), bold, light gray 5 fill
        header_row = current_row
        ws1.row_dimensions[header_row].height = 137.52
        for col_idx, header in enumerate(COMPLIANCE_HEADERS, start=2):
            c = ws1.cell(row=header_row, column=col_idx, value=header)
            c.font = Font(name=FONT_NAME, size=9, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", textRotation=90, wrap_text=True)
            c.fill = LIGHT_GRAY_5_FILL

        # Bold outer border around B2:J2
        apply_outer_border(ws1, header_row, header_row, 2, 10)
        current_row += 1

        start_data_row = current_row
        for d in dates:
            if dataset_type == "shift_a":
                st = shift_a[d]
            elif dataset_type == "shift_b":
                st = shift_b[d]
            else:  # combined
                sa = shift_a[d]
                sb = shift_b[d]
                st = {
                    "weighed": sa["weighed"] + sb["weighed"],
                    "warned_charged": sa["warned_charged"] + sb["warned_charged"],
                    "charged": sa["charged"] + sb["charged"],
                    "warned": sa["warned"] + sb["warned"],
                    "cases_cleared": sa["cases_cleared"] + sb["cases_cleared"],
                    "transgressions": sa["transgressions"] + sb["transgressions"],
                    "exempted": sa["exempted"] + sb["exempted"],
                    "manually_weighed": sa["manually_weighed"] + sb["manually_weighed"],
                }

            ws1.row_dimensions[current_row].height = 15.84

            # Date (Col B) - text in bold, centered
            dt_obj = datetime.strptime(d, "%Y-%m-%d")
            c_date = ws1.cell(row=current_row, column=2, value=dt_obj)
            c_date.font = Font(name=FONT_NAME, size=10, bold=True)
            c_date.alignment = align_center
            c_date.border = THIN_BORDER
            c_date.number_format = "d-mmm-yy"

            # Values (Cols C to J) - calibri, aligned and justified center
            vals = [
                st["weighed"],
                st["warned_charged"],
                st["charged"],
                st["warned"],
                st["cases_cleared"],
                st["transgressions"],
                st["exempted"],
                st["manually_weighed"],
            ]
            for col_idx, v in enumerate(vals, start=3):
                cell = ws1.cell(row=current_row, column=col_idx, value=v)
                cell.font = Font(name=FONT_NAME, size=10, bold=False)
                cell.alignment = align_center
                cell.border = THIN_BORDER
                cell.number_format = "#,##0"

            current_row += 1

        end_data_row = current_row - 1

        # TOTAL Row (B34:J34) - bold text, light gray 5 fill, bold border
        total_row = current_row
        ws1.row_dimensions[total_row].height = 20
        t_label = ws1.cell(row=total_row, column=2, value="TOTAL")
        t_label.font = Font(name=FONT_NAME, size=10, bold=True)
        t_label.alignment = align_center
        t_label.fill = LIGHT_GRAY_5_FILL

        for col_idx in range(3, 11):
            col_letter = get_column_letter(col_idx)
            formula = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
            c_tot = ws1.cell(row=total_row, column=col_idx, value=formula)
            c_tot.font = Font(name=FONT_NAME, size=10, bold=True)
            c_tot.alignment = align_center
            c_tot.fill = LIGHT_GRAY_5_FILL
            c_tot.number_format = "#,##0"

        # Bold outer border around B:J for TOTAL row
        apply_outer_border(ws1, total_row, total_row, 2, 10)

        current_row += 3  # Spacers

    write_compliance_block("MOBILE SHIFT A", "shift_a")
    write_compliance_block("MOBILE SHIFT B", "shift_b")
    write_compliance_block("MOBILE TOTAL", "combined")

    # Column dimensions for Sheet 1
    ws1.column_dimensions["A"].width = 3.0
    ws1.column_dimensions["B"].width = 15.0  # 1.27''
    ws1.column_dimensions["C"].width = 9.0   # 0.76''
    ws1.column_dimensions["D"].width = 9.0   # 0.76''
    ws1.column_dimensions["E"].width = 15.0  # 1.27''
    ws1.column_dimensions["F"].width = 9.0   # 0.76''
    ws1.column_dimensions["G"].width = 9.0   # 0.76''
    ws1.column_dimensions["H"].width = 9.0   # 0.76''
    ws1.column_dimensions["I"].width = 9.0   # 0.76''
    ws1.column_dimensions["J"].width = 9.0   # 0.76''

    # =========================================================================
    # SHEET 2: MOBILE PER HOUR
    # =========================================================================
    ws2 = wb.create_sheet(title="MOBILE PER HOUR")
    ws2.views.sheetView[0].showGridLines = True

    h_row = 1

    def write_hourly_block(title: str, dataset_type: str):
        nonlocal h_row

        # Title Row: merged A to Z
        ws2.row_dimensions[h_row].height = 20
        ws2.merge_cells(start_row=h_row, start_column=1, end_row=h_row, end_column=26)
        t_cell = ws2.cell(row=h_row, column=1, value=title)
        t_cell.font = font_title
        t_cell.alignment = align_center
        for col in range(1, 27):
            ws2.cell(row=h_row, column=col).border = THIN_BORDER
        h_row += 1

        # Header Row with Vertical Rotated Hours
        header_row = h_row
        ws2.row_dimensions[header_row].height = 70
        h_date = ws2.cell(row=header_row, column=1, value="Date")
        h_date.font = font_header_8
        h_date.alignment = align_center_wrap

        for col_idx, h_label in enumerate(HOURLY_DISPLAY_HEADERS, start=2):
            c = ws2.cell(row=header_row, column=col_idx, value=h_label)
            c.font = font_header_8
            c.alignment = align_vertical_header

        c_tot = ws2.cell(row=header_row, column=26, value="Total Weighed per day")
        c_tot.font = font_header_8
        c_tot.alignment = align_center_wrap

        # Apply outer bold border surrounding all columns of header row from A to Z
        apply_outer_border(ws2, header_row, header_row, 1, 26)
        h_row += 1

        start_data_row = h_row
        for d in dates:
            if dataset_type == "shift_a":
                hrly = shift_a[d]["hourly"]
            elif dataset_type == "shift_b":
                hrly = shift_b[d]["hourly"]
            else:
                sa = shift_a[d]["hourly"]
                sb = shift_b[d]["hourly"]
                hrly = {h: sa[h] + sb[h] for h in HOURS}

            ws2.row_dimensions[h_row].height = 18

            # Date (Col A) - font size 11 and bold text, center aligned
            dt_obj = datetime.strptime(d, "%Y-%m-%d")
            c_date = ws2.cell(row=h_row, column=1, value=dt_obj)
            c_date.font = font_date_11
            c_date.alignment = align_center
            c_date.border = THIN_BORDER
            c_date.number_format = "d-mmm-yy"

            # 24 Hours (Col B to Y) - font size 11, not bold, center aligned
            for col_idx, h_key in enumerate(HOURS, start=2):
                cell = ws2.cell(row=h_row, column=col_idx, value=hrly.get(h_key, 0))
                cell.font = font_data_11
                cell.alignment = align_center
                cell.border = THIN_BORDER
                cell.number_format = "#,##0"

            # Total Weighed per day (Col Z) - font size 11, not bold, center aligned
            day_sum_formula = f"=SUM(B{h_row}:Y{h_row})"
            cell_sum = ws2.cell(row=h_row, column=26, value=day_sum_formula)
            cell_sum.font = font_data_11
            cell_sum.alignment = align_center
            cell_sum.border = THIN_BORDER
            cell_sum.number_format = "#,##0"

            h_row += 1

        end_data_row = h_row - 1

        # Total Weighed per Hour Row (Row 34 for 31-day month)
        total_row = h_row
        ws2.row_dimensions[total_row].height = 28
        t_label = ws2.cell(row=total_row, column=1, value="Total Weighed per Hour")
        t_label.font = font_summary_11
        t_label.alignment = align_left_middle_wrap

        for col_idx in range(2, 27):
            col_letter = get_column_letter(col_idx)
            formula = f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
            c_tot = ws2.cell(row=total_row, column=col_idx, value=formula)
            c_tot.font = font_summary_11
            c_tot.alignment = align_center
            c_tot.number_format = "#,##0"

        h_row += 1

        # Average Row (Row 35 for 31-day month) - whole numbers
        avg_row = h_row
        ws2.row_dimensions[avg_row].height = 20
        avg_label = ws2.cell(row=avg_row, column=1, value="Average")
        avg_label.font = font_summary_11
        avg_label.alignment = align_left_middle_wrap

        for col_idx in range(2, 27):
            col_letter = get_column_letter(col_idx)
            formula = f"=ROUND(AVERAGE({col_letter}{start_data_row}:{col_letter}{end_data_row}), 0)"
            c_avg = ws2.cell(row=avg_row, column=col_idx, value=formula)
            c_avg.font = font_summary_11
            c_avg.alignment = align_center
            c_avg.number_format = "#,##0"

        # Apply outer bold border surrounding total and average rows (A to Z)
        apply_outer_border(ws2, total_row, avg_row, 1, 26)

        h_row += 2  # 1 blank spacer row between tables

    write_hourly_block("MOBILE SHIFT A", "shift_a")
    write_hourly_block("MOBILE SHIFT B", "shift_b")
    write_hourly_block("TOTAL PER HOUR", "combined")

    # Column widths for Sheet 2
    ws2.column_dimensions["A"].width = 13.0
    for col_idx in range(2, 26):
        ws2.column_dimensions[get_column_letter(col_idx)].width = 4.8
    ws2.column_dimensions["Z"].width = 14.0

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
